#!/usr/bin/env python3
"""
Lead Discovery Tool v3.0
========================
Search the web for companies and events, get actionable leads in Excel.
Uses DuckDuckGo for search, optional page visit for email enrichment.

Run: python3 main.py
"""

from __future__ import annotations

import asyncio
import csv
import os
import platform
import re
import sqlite3
import subprocess
import sys
import time
from concurrent.futures import ThreadPoolExecutor
from datetime import datetime
from pathlib import Path
from urllib.parse import urlparse

# ── 0. Dependency check + install ──────────────────────────────────────────
MISSING = []
try: from rich.console import Console
except ImportError: MISSING.append("rich")
try: from openpyxl import Workbook; from openpyxl.styles import Font, PatternFill, Alignment; from openpyxl.utils import get_column_letter
except ImportError: MISSING.append("openpyxl")
try: from ddgs import DDGS
except ImportError: MISSING.append("ddgs")
try: import aiohttp
except ImportError: MISSING.append("aiohttp")
try: from bs4 import BeautifulSoup
except ImportError: MISSING.append("beautifulsoup4")

if MISSING:
    in_venv = hasattr(sys, "base_prefix") and sys.base_prefix != sys.prefix or "VIRTUAL_ENV" in os.environ
    if not in_venv:
        print("Please create a virtual environment first:")
        print("  python3 -m venv venv && source venv/bin/activate && pip install rich openpyxl ddgs aiohttp beautifulsoup4")
        sys.exit(1)
    print(f"Installing: {MISSING} ...")
    subprocess.check_call([sys.executable, "-m", "pip", "install", "-q"] + MISSING)
    # re-import after install
    from rich.console import Console
    from openpyxl import Workbook; from openpyxl.styles import Font, PatternFill, Alignment; from openpyxl.utils import get_column_letter
    from ddgs import DDGS
    import aiohttp
    from bs4 import BeautifulSoup

console = Console()

# ── 1. Project root + directories ──────────────────────────────────────────
ROOT = Path(__file__).resolve().parent
os.chdir(ROOT)
(ROOT / "exports").mkdir(exist_ok=True)
(ROOT / "data").mkdir(exist_ok=True)

USER_AGENT = "Mozilla/5.0 (Macintosh; Intel Mac OS X 10_15_7) AppleWebKit/537.36 Chrome/125 Safari/537.36"

# ── 2. Search ──────────────────────────────────────────────────────────────

def search(query: str, max_results: int = 10) -> list[dict]:
    """Search DDG, return [{title, url, description}]."""
    try:
        with DDGS() as d:
            raw = list(d.text(query, max_results=max_results))
        return [
            {"title": _clean_title(r["title"]), "url": r["href"], "description": r.get("body", "")}
            for r in raw
        ]
    except Exception as e:
        console.print(f"  [yellow]Search error: {e}[/]")
        return []

def _clean_title(title: str) -> str:
    for s in [" - Wikipedia", " | LinkedIn", " - Crunchbase", " - Home", " | Facebook", " - YouTube", " · Instagram"]:
        if s in title:
            title = title.split(s)[0]
    if " | " in title and len(title.split(" | ")[0]) > 8:
        title = title.split(" | ")[0]
    return title.strip()[:120]

# ── 3. Page enrichment (lightweight) ───────────────────────────────────────

async def enrich(url: str, session: aiohttp.ClientSession) -> dict:
    """Visit page, extract emails + hiring signal. Returns {emails, hiring}."""
    result = {"emails": [], "hiring": False}
    try:
        async with session.get(url, headers={"User-Agent": USER_AGENT}, timeout=aiohttp.ClientTimeout(total=8), allow_redirects=True) as resp:
            if resp.status != 200:
                return result
            ct = resp.headers.get("Content-Type", "")
            if "text/html" not in ct:
                return result
            html = await resp.text()
    except Exception:
        return result

    soup = BeautifulSoup(html, "html.parser")
    emails = re.findall(r"[a-zA-Z0-9_.+-]+@[a-zA-Z0-9-]+\.[a-zA-Z]{2,}", html)
    generic = {"support", "info", "admin", "contact", "sales", "billing", "hello", "noreply", "no-reply", "webmaster", "postmaster"}
    unique = []
    for e in emails:
        pfx = e.split("@")[0].lower()
        if pfx not in generic and e not in unique and not re.search(r"[;&<>#]", e) and not e.endswith(".") and "u003" not in e and e.split(".")[-1].isalpha():
            unique.append(e)
            if len(unique) >= 3:
                break
    result["emails"] = unique

    career_words = {"career", "hiring", "job", "join", "vacanc", "work with us"}
    for a in soup.find_all("a", href=True):
        t = a.get_text(strip=True).lower()
        h = a["href"].lower()
        if any(w in t for w in career_words) or any(w in h for w in career_words):
            result["hiring"] = True
            break
    return result

# ── 4. Company discovery ───────────────────────────────────────────────────

async def discover_companies(keyword: str, count: int) -> list[dict]:
    """Search DDG with 3 targeted queries, deduplicate, optionally enrich."""
    queries = [
        f'{keyword} company',
        f'{keyword} services',
        f'{keyword} contact',
    ]
    per = max(5, count // 3 + 2)
    seen = set()
    leads = []

    with ThreadPoolExecutor(max_workers=1) as ex:
        for q in queries:
            console.print(f"  🔍 [dim]{q}[/]")
            results = await asyncio.get_event_loop().run_in_executor(ex, lambda q=q: search(q, per))
            await asyncio.sleep(1.5)
            for r in results:
                if r["url"] not in seen and len(leads) < count:
                    seen.add(r["url"])
                    leads.append(r)
            if len(leads) >= count:
                break

    console.print(f"  📋 [bold]{len(leads)} unique leads found[/]\n")

    async with aiohttp.ClientSession(timeout=aiohttp.ClientTimeout(total=25)) as session:
        companies = []
        for i, lead in enumerate(leads, 1):
            url = lead["url"]
            desc = lead["description"]
            hiring = "No"
            emails = []

            # Enrich first half of results only (saves time)
            if i <= max(3, count // 2):
                console.print(f"  [{i}/{len(leads)}] Enriching: {lead['title'][:45]}...")
                info = await enrich(url, session)
                emails = info["emails"]
                hiring = "Yes" if info["hiring"] else _hiring_snippet(desc)
            else:
                console.print(f"  [{i}/{len(leads)}] {lead['title'][:45]}...")
                hiring = _hiring_snippet(desc)

            companies.append({
                "name": lead["title"],
                "website": url,
                "industry": keyword,
                "description": desc,
                "emails": ", ".join(emails),
                "hiring": hiring,
                "size": _guess_size(desc),
            })
            await asyncio.sleep(2)

        return companies

def _hiring_snippet(text: str) -> str:
    if any(s in text.lower() for s in ["hiring", "career", "job opening", "join our team", "we're hiring"]):
        return "Yes"
    return "No"

def _guess_size(text: str) -> str:
    t = text.lower()
    if any(w in t for w in ["enterprise", "global", "fortune"]): return "Large"
    if any(w in t for w in ["startup", "small business", "boutique", "early-stage"]): return "Small"
    return "Mid"

# ── 5. Event discovery ─────────────────────────────────────────────────────

async def discover_events(city: str, count: int) -> list[dict]:
    queries = [
        f"{city} networking events mixer",
        f"{city} business meetup conference",
    ]
    per = max(5, count // 2 + 2)
    seen = set()
    leads = []

    with ThreadPoolExecutor(max_workers=1) as ex:
        for q in queries:
            console.print(f"  🎉 [dim]{q}[/]")
            results = await asyncio.get_event_loop().run_in_executor(ex, lambda q=q: search(q, per))
            await asyncio.sleep(1.5)
            for r in results:
                if r["url"] not in seen and len(leads) < count:
                    seen.add(r["url"])
                    leads.append(r)
            if len(leads) >= count:
                break

    console.print(f"  📋 [bold]{len(leads)} events found[/]\n")

    refresh_words = ["refreshment", "drinks", "buffet", "dinner", "happy hour", "cocktail", "catered", "open bar", "coffee", "wine", "breakfast", "lunch", "canap"]
    events = []
    for lead in leads:
        t = (lead["title"] + " " + lead["description"]).lower()
        signal = "None"
        for w in refresh_words:
            if w in t:
                signal = f"Mentions '{w}'"
                break
        events.append({
            "name": lead["title"],
            "date": _extract_date(lead["description"]),
            "venue": city,
            "url": lead["url"],
            "refreshments": signal,
            "description": lead["description"],
        })
    return events

def _extract_date(text: str) -> str:
    m = re.search(r"\b\d{4}-\d{2}-\d{2}\b", text)
    if m: return m.group()
    m = re.search(r"\b(?:Jan|Feb|Mar|Apr|May|Jun|Jul|Aug|Sep|Oct|Nov|Dec)[a-z]* \d{1,2},? \d{4}\b", text, re.I)
    if m: return m.group()
    return "TBA"

# ── 6. Excel export ────────────────────────────────────────────────────────

def export(companies=None, events=None) -> str:
    ts = datetime.now().strftime("%Y%m%d_%H%M%S")
    wb = Workbook()
    wb.remove(wb.active)
    hf = Font(bold=True, color="FFFFFF", size=11)
    hfill = PatternFill(start_color="4F81BD", end_color="4F81BD", fill_type="solid")

    def style(ws):
        for c in ws[1]: c.font, c.fill = hf, hfill
        for col in ws.columns:
            mx = 10
            for cell in col:
                if cell.value: mx = max(mx, min(len(str(cell.value)), 55))
            ws.column_dimensions[get_column_letter(col[0].column)].width = mx + 2

    # Overview
    ws = wb.create_sheet("Overview")
    ws.append(["Metric", "Value"])
    ws.append(["Generated", datetime.now().strftime("%Y-%m-%d %H:%M:%S")])
    ws.append(["Companies", str(len(companies)) if companies else "0"])
    ws.append(["Events", str(len(events)) if events else "0"])
    ws.append(["", ""])
    ws.append(["👇 Check the Companies, Events & Contacts tabs below", ""])
    style(ws)

    # Companies
    if companies:
        ws = wb.create_sheet("Companies")
        ws.append(["Company", "Website", "Industry", "Size", "Hiring", "Description", "Emails"])
        for c in companies:
            ws.append([c.get(k, "") for k in ["name", "website", "industry", "size", "hiring", "description", "emails"]])
        style(ws)

        # Contacts (non-empty emails only)
        contacts = [(c["name"], c["emails"]) for c in companies if c.get("emails")]
        if contacts:
            ws = wb.create_sheet("Contacts")
            ws.append(["Company", "Emails"])
            for name, emails in contacts:
                ws.append([name, emails])
            style(ws)

    # Events
    if events:
        ws = wb.create_sheet("Events")
        ws.append(["Event", "Date", "City", "URL", "Refreshments", "Description"])
        for e in events:
            ws.append([e.get(k, "") for k in ["name", "date", "venue", "url", "refreshments", "description"]])
        style(ws)

    path = str(ROOT / "exports" / f"Leads_{ts}.xlsx")
    wb.save(path)
    return path

# ── 7. CLI ─────────────────────────────────────────────────────────────────

def show_banner():
    console.print()
    console.print("  [bold cyan]🔍 LEAD DISCOVERY TOOL[/]")
    console.print("  [dim]Search the web → find companies & events → export to Excel[/]")
    console.print()

async def main():
    show_banner()
    while True:
        console.print()
        console.print("  [1] [bold]Find Companies[/] by keyword")
        console.print("  [2] [bold]Find Events[/] by city")
        console.print("  [3] Exit")
        console.print()

        c = input("  Choice [1-3]: ").strip()
        if c == "1":
            kw = input("  Industry / keyword: ").strip()
            n = input("  Max results [10]: ").strip()
            n = int(n) if n.isdigit() else 10
            console.print()
            companies = await discover_companies(kw, n)
            if companies:
                path = export(companies=companies)
                console.print(f"\n  ✅ [bold green]Saved: {path}[/]")
                _open(path)
            else:
                console.print("  [yellow]No results found.[/]")
            input("\n  Press Enter to continue...")

        elif c == "2":
            city = input("  City: ").strip()
            n = input("  Max results [10]: ").strip()
            n = int(n) if n.isdigit() else 10
            console.print()
            events = await discover_events(city, n)
            if events:
                path = export(events=events)
                console.print(f"\n  ✅ [bold green]Saved: {path}[/]")
                _open(path)
            else:
                console.print("  [yellow]No events found.[/]")
            input("\n  Press Enter to continue...")

        elif c == "3":
            console.print("  👋 Bye!\n")
            break

def _open(path: str):
    try:
        s = platform.system()
        if s == "Darwin": subprocess.run(["open", path])
        elif s == "Windows": os.startfile(path)
        else: subprocess.run(["xdg-open", path])
    except Exception:
        pass

if __name__ == "__main__":
    try:
        asyncio.run(main())
    except KeyboardInterrupt:
        print("\n  👋 Bye!")
